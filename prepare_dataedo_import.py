import json
from pathlib import Path
from openpyxl import Workbook


# File input dan output
INPUT_FILE = Path("output/metadata_clean.json")
OUTPUT_FILE = Path("output/dataedo_metadata_import.xlsx")


# Baca metadata JSON
with open(INPUT_FILE, "r", encoding="utf-8") as f:
    metadata = json.load(f)


# Buat Excel
wb = Workbook()


# ============================================================
# 1. TABLES
# ============================================================

ws = wb.active
ws.title = "Tables"

ws.append([
    "Database",
    "Schema",
    "Table",
    "ObjectType"
])

database = metadata.get("database") or ""
schema = metadata.get("schema") or ""
view_name = metadata.get("name") or ""

# Target view
ws.append([
    database,
    schema,
    view_name,
    "VIEW"
])

# Source tables
for source in metadata.get("source_objects", []):
    ws.append([
        source.get("database", ""),
        source.get("schema", ""),
        source.get("table", ""),
        "TABLE"
    ])


# ============================================================
# 2. COLUMNS
# ============================================================

ws_columns = wb.create_sheet("Columns")

ws_columns.append([
    "Database",
    "Schema",
    "Table",
    "Column",
    "DataType",
    "Description"
])

for column in metadata.get("columns", []):
    ws_columns.append([
        database,
        schema,
        view_name,
        column.get("name", ""),
        "",
        column.get("expression", "")
    ])


# ============================================================
# 3. RELATIONSHIPS
# ============================================================

ws_rel = wb.create_sheet("Relationships")

ws_rel.append([
    "Source Table",
    "Source Column",
    "Target Table",
    "Target Column"
])

for rel in metadata.get("relationships", []):

    source_table = rel.get("source_object", "")
    target_table = rel.get("target_object", "")

    column_pairs = rel.get("column_pairs", [])

    for pair in column_pairs:

        if isinstance(pair, dict):
            source_column = (
                pair.get("source_column")
                or pair.get("source")
                or ""
            )

            target_column = (
                pair.get("target_column")
                or pair.get("target")
                or ""
            )
        else:
            source_column = ""
            target_column = ""

        ws_rel.append([
            source_table,
            source_column,
            target_table,
            target_column
        ])


# ============================================================
# 4. LINEAGE
# ============================================================

ws_lin = wb.create_sheet("Lineage")

ws_lin.append([
    "Source",
    "Processor",
    "Target",
    "Source Column",
    "Target Column",
    "Transformation"
])

for column in metadata.get("columns", []):

    target_column = column.get("name", "")
    expression = column.get("expression", "")

    dependencies = column.get("dependencies", [])

    for dependency in dependencies:

        if isinstance(dependency, dict):

            source = (
                dependency.get("source_object")
                or dependency.get("object")
                or dependency.get("source")
                or ""
            )

            source_column = (
                dependency.get("column")
                or dependency.get("source_column")
                or ""
            )

        else:
            source = str(dependency)
            source_column = ""

        ws_lin.append([
            source,
            view_name,
            view_name,
            source_column,
            target_column,
            expression
        ])


# ============================================================
# FORMATTING
# ============================================================

for worksheet in wb.worksheets:

    worksheet.freeze_panes = "A2"

    # Header bold
    for cell in worksheet[1]:
        cell.font = cell.font.copy(bold=True)

    # Auto width
    for column_cells in worksheet.columns:

        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12),
            60
        )


# ============================================================
# SAVE
# ============================================================

wb.save(OUTPUT_FILE)

print("")
print("========================================")
print("Dataedo metadata file created!")
print("========================================")
print(f"Input : {INPUT_FILE}")
print(f"Output: {OUTPUT_FILE}")
print("")
print("Sheets:")
print("1. Tables")
print("2. Columns")
print("3. Relationships")
print("4. Lineage")